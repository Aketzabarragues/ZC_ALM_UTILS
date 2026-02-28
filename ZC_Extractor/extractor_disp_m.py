import pandas as pd
import os
from openpyxl import load_workbook
import openpyxl.utils as utils

def extraer(ruta_excel, folder_path):
    try:
        # 1. Cargar el Excel y localizar la tabla
        wb = load_workbook(ruta_excel, data_only=True)
        ws = wb["DISP_M"]
        tabla = next((t for t in ws.tables.values() if t.name == "Tabla_Disp_M"), None)
        
        if not tabla:
            return False, "No se encontró Tabla_Disp_M"

        # 2. Leer los datos basándose en el rango de la tabla
        min_col, min_row, max_col, max_row = utils.cell.range_boundaries(tabla.ref)
        df = pd.read_excel(ruta_excel, sheet_name="DISP_M",
                           skiprows=min_row - 1,
                           usecols=list(range(min_col - 1, max_col)),
                           nrows=max_row - min_row + 1)

        # 3. Limpieza de datos
        # Eliminamos filas sin UID y rellenamos vacíos para evitar errores en XML
        df = df.dropna(subset=['UID']).fillna('')

        # --- GUARDAR ARCHIVO XML ---
        output_file = os.path.join(folder_path, "disp_m.xml")
        
        # Exportamos a XML con una estructura jerárquica clara
        # root_name: El nodo principal
        # row_name: El nombre de cada objeto
        # attr_cols=False: Para que use etiquetas <Campo> en lugar de atributos
        df.to_xml(output_file, 
                  index=False, 
                  root_name="Dispositivos", 
                  row_name="Dispositivo", 
                  attr_cols=False, 
                  xml_declaration=True, 
                  encoding='utf-8')

        return True, f"Exportado exitosamente a {output_file}"

    except Exception as e:
        return False, str(e)